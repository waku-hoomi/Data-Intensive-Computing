"""Render the Week 4 design and evaluation Markdown reports to .docx.

The Markdown files under docs/ are the source of truth; this converts the
subset of Markdown they use (headings, bold, inline code, bullets, pipe
tables, explicit page breaks) into a Word document. Word is used rather than
a PDF writer so the group can open, adjust and export the reports themselves
in whatever tool the course submission expects.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
REPORTS = ROOT / "reports"

NAVY = RGBColor(0x18, 0x33, 0x4D)
TEAL = RGBColor(0x17, 0x6B, 0x70)
BODY_GREY = RGBColor(0x1D, 0x28, 0x30)

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def _add_runs(paragraph, text: str, *, base_bold: bool = False, color: RGBColor | None = None) -> None:
    """Write one line into a paragraph, honouring **bold** and `code` spans."""
    for piece in INLINE_RE.split(text):
        if not piece:
            continue
        bold = base_bold
        italic = False
        font = None
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            piece, bold = piece[2:-2], True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            piece, font = piece[1:-1], "Consolas"
        run = paragraph.add_run(piece)
        run.bold = bold
        run.italic = italic
        if font:
            run.font.name = font
            run.font.size = Pt(9.5)
        if color is not None:
            run.font.color.rgb = color


def _style_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY_GREY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08


def _parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for index, line in enumerate(lines):
        if index == 1:  # the | --- | separator
            continue
        rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
    return rows


def markdown_to_docx(markdown_path: Path, output_path: Path, title: str) -> None:
    document = Document()
    _style_document(document)
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0

    while index < len(lines):
        line = lines[index]

        if line.startswith("# "):
            heading = document.add_heading(level=0)
            _add_runs(heading, line[2:], base_bold=True, color=NAVY)
            index += 1
            continue

        if line.startswith("## "):
            heading = document.add_heading(level=1)
            _add_runs(heading, line[3:], base_bold=True, color=TEAL)
            index += 1
            continue

        if line.strip() == "<!-- PAGEBREAK -->":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and lines[index + 1].startswith("| ---"):
            table_lines = [line, lines[index + 1]]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = _parse_table(table_lines)
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    target = table.cell(r, c).paragraphs[0]
                    _add_runs(target, cell, base_bold=(r == 0))
                    for run in target.runs:
                        run.font.size = Pt(9)
            document.add_paragraph()
            continue

        if line.startswith("- "):
            while index < len(lines) and lines[index].startswith("- "):
                item = lines[index][2:]
                paragraph = document.add_paragraph(style="List Bullet")
                _add_runs(paragraph, item)
                index += 1
            continue

        if not line.strip():
            index += 1
            continue

        # Body paragraph: join consecutive non-blank, non-structural lines.
        block = [line]
        index += 1
        while (index < len(lines) and lines[index].strip()
               and not lines[index].startswith(("#", "|", "- "))
               and lines[index].strip() != "<!-- PAGEBREAK -->"):
            block.append(lines[index])
            index += 1
        paragraph = document.add_paragraph()
        _add_runs(paragraph, " ".join(part.strip() for part in block))

    core = document.core_properties
    core.title = title
    core.author = "ID2221 group"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def main() -> None:
    markdown_to_docx(DOCS / "Week4_Design_Report.md",
                     REPORTS / "Week4_Design_Report.docx", "Week 4 Design Report")
    markdown_to_docx(DOCS / "Week4_Evaluation_Report.md",
                     REPORTS / "Week4_Evaluation_Report.docx", "Week 4 Evaluation Report")
    print("design_docx:", REPORTS / "Week4_Design_Report.docx")
    print("evaluation_docx:", REPORTS / "Week4_Evaluation_Report.docx")


if __name__ == "__main__":
    main()
